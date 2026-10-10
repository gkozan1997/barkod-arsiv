import os
import openpyxl
import csv

def normalize_header(s):
    if not s:
        return ""
    mapping = {
        'ç': 'c', 'ğ': 'g', 'ı': 'i', 'ö': 'o', 'ş': 's', 'ü': 'u',
        'Ç': 'c', 'Ğ': 'g', 'I': 'i', 'İ': 'i', 'Ö': 'o', 'Ş': 's', 'Ü': 'u'
    }
    s = str(s).lower()
    for k, v in mapping.items():
        s = s.replace(k, v)
    return "".join(c for c in s if c.isalnum())

def find_column_index(headers, possible_names):
    clean_headers = [normalize_header(h) for h in headers]
    for name in possible_names:
        n_clean = normalize_header(name)
        # 1. Tam eşleşme önceliği
        for idx, h_clean in enumerate(clean_headers):
            if h_clean == n_clean:
                return idx
        # 2. İçerme önceliği
        for idx, h_clean in enumerate(clean_headers):
            if n_clean in h_clean:
                return idx
    return None

def parse_order_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    rows = []

    if ext in ['.xlsx', '.xlsm']:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(list(row))
    elif ext == '.csv':
        with open(file_path, 'r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.reader(f)
            for row in reader:
                if any(row):
                    rows.append(row)
    elif ext == '.xls':
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            sheet = wb.sheet_by_index(0)
            for r in range(sheet.nrows):
                rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        except Exception as e:
            print(f"XLS okuma hatası: {e}")
    else:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append(list(row))
        except Exception:
            pass

    if not rows:
        return []

    # Başlık satırını bul
    header_idx = -1
    for i in range(min(6, len(rows))):
        row_str = " ".join([normalize_header(c) for c in rows[i] if c is not None])
        if any(keyword in row_str for keyword in ['urun', 'siparis', 'adet', 'miktar', 'pazaryeri']):
            header_idx = i
            break

    if header_idx == -1:
        header_idx = 0

    headers = [str(c) if c is not None else "" for c in rows[header_idx]]
    data_rows = rows[header_idx + 1:]

    col_siparis = find_column_index(headers, ['siparisno', 'siparisnumarasi', 'paketno', 'ordernumber', 'orderno'])
    col_paket = find_column_index(headers, ['paketno'])
    col_musteri = find_column_index(headers, ['sevkmusteri', 'faturamusteri', 'musteriadi', 'alici', 'adsoyad', 'customer'])
    col_urun = find_column_index(headers, ['urun', 'urunadi', 'productname', 'ilanbasligi', 'itemname'])
    col_adet = find_column_index(headers, ['adet', 'miktar', 'siparisadedi', 'quantity', 'qty'])
    col_pazar = find_column_index(headers, ['pazaryeri', 'pazaryeriadi', 'platform', 'marketplace'])
    col_stok = find_column_index(headers, ['stokkodu', 'sku'])
    col_gtin = find_column_index(headers, ['gtin', 'gtinbarkod', 'barkod'])

    filename_lower = os.path.basename(file_path).lower()
    default_pazar = "Diğer"
    if "trendyol" in filename_lower:
        default_pazar = "Trendyol"
    elif "hepsiburada" in filename_lower or "hb" in filename_lower:
        default_pazar = "Hepsiburada"
    elif "cicek" in filename_lower:
        default_pazar = "Çiçeksepeti"
    elif "pazarama" in filename_lower:
        default_pazar = "Pazarama"
    elif "amazon" in filename_lower:
        default_pazar = "Amazon"
    elif "n11" in filename_lower:
        default_pazar = "N11"

    orders_map = {}
    for r_idx, r in enumerate(data_rows):
        if not r or not any(r):
            continue

        siparis_no = str(r[col_siparis]).strip() if col_siparis is not None and col_siparis < len(r) and r[col_siparis] is not None else f"SIP-{r_idx+1}"
        paket_no = str(r[col_paket]).strip() if col_paket is not None and col_paket < len(r) and r[col_paket] is not None else ""
        musteri_adi = str(r[col_musteri]).strip() if col_musteri is not None and col_musteri < len(r) and r[col_musteri] is not None else "Belirtilmemiş"
        urun_adi = str(r[col_urun]).strip() if col_urun is not None and col_urun < len(r) and r[col_urun] is not None else "Ürün"
        pazaryeri = str(r[col_pazar]).strip() if col_pazar is not None and col_pazar < len(r) and r[col_pazar] is not None else default_pazar
        stok_kod = str(r[col_stok]).strip() if col_stok is not None and col_stok < len(r) and r[col_stok] is not None else ""
        gtin = str(r[col_gtin]).strip() if col_gtin is not None and col_gtin < len(r) and r[col_gtin] is not None else ""

        adet = 1
        if col_adet is not None and col_adet < len(r) and r[col_adet] is not None:
            try:
                adet = int(float(str(r[col_adet]).replace(',', '.')))
                if adet < 1:
                    adet = 1
            except Exception:
                adet = 1

        if siparis_no not in orders_map:
            orders_map[siparis_no] = {
                'siparis_no': siparis_no,
                'paket_no': paket_no,
                'musteri_adi': musteri_adi,
                'pazaryeri': pazaryeri,
                'urunler': [],
                'tamamlandi': False
            }

        item_id = f"item_{len(orders_map[siparis_no]['urunler']) + 1}"
        orders_map[siparis_no]['urunler'].append({
            'item_id': item_id,
            'urun_adi': urun_adi,
            'stok_kod': stok_kod,
            'gtin': gtin,
            'adet': adet,
            'okunan_adet': 0,
            'seri_nolar': [],
            'tamamlandi': False
        })

    return list(orders_map.values())
